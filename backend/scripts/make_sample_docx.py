"""Generate a sample DOCX for the browser upload journey.

Run from ``backend/`` with the shared venv python; the destination path is
``argv[1]``. One ``Heading 1`` plus one paragraph covers both element types
the stub DOCX parser produces, mirroring the committed ``tiny.pdf`` fixture
so the journey exercises the heading and paragraph paths of the pipeline.
"""

import sys

from docx import Document


def main() -> None:
    """Write a minimal DOCX with one heading and one paragraph to ``argv[1]``."""

    if len(sys.argv) != 2:
        raise SystemExit("usage: make_sample_docx.py <destination.docx>")
    document = Document()
    document.add_heading("Sample heading", level=1)
    document.add_paragraph("Sample paragraph text for the upload journey.")
    document.save(sys.argv[1])


if __name__ == "__main__":
    main()
