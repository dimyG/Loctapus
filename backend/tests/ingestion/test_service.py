"""Service tests for the ingest task, the embed task, and the dimension guard."""

import hashlib
import io
from dataclasses import replace
from uuid import UUID, uuid4

import pytest
from docx import Document as DocxDocument
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.ingestion.models import Chunk, Document, IngestionRun
from app.ingestion.service import (
    canonical_document_json,
    embed_document,
    ingest_source_item,
)
from app.outbox.events import DOCUMENT_READY_FOR_EMBEDDING
from app.outbox.models import OutboxEvent
from app.providers import get_embedding_provider, get_parser
from app.providers.base import EmbeddingProvider
from app.providers.stub_embeddings import STUB_EMBEDDING_DIM
from app.sources.models import Source, SourceItem
from app.storage.service import StorageService
from tests.factories import SeededUpload, seed_project, seed_upload
from tests.helpers import SettingsForTests


@pytest.fixture
def storage(*, settings: Settings) -> StorageService:
    """Return the real storage service pointed at the compose Garage with the dev bucket."""

    return StorageService(settings=settings)


class _CountingEmbeddingProvider:
    """Embedding provider wrapper that records the size of every embed_documents call."""

    dimension: int

    def __init__(self, *, inner: EmbeddingProvider) -> None:
        """Wrap ``inner`` and start with an empty call-size record."""

        self._inner = inner
        self.dimension = inner.dimension
        self.call_sizes: list[int] = []

    def embed_documents(self, *, texts: list[str]) -> list[list[float]]:
        """Delegate to the wrapped provider and record this call's batch size."""

        self.call_sizes.append(len(texts))
        return self._inner.embed_documents(texts=texts)

    def embed_query(self, *, text: str) -> list[float]:
        """Delegate a single-text embedding to the wrapped provider."""

        return self._inner.embed_query(text=text)


def _make_docx(*, heading: str | None, paragraphs: list[str]) -> bytes:
    """Build an in-memory DOCX carrying one optional heading and the given paragraphs."""

    buffer = io.BytesIO()
    document = DocxDocument()
    if heading is not None:
        document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(buffer)
    return buffer.getvalue()


def _seed(
    *,
    db: Session,
    storage: StorageService,
    filename: str,
    content: bytes,
) -> tuple[UUID, SeededUpload]:
    """Seed one upload with its bytes in Garage: a project, its rows, and the stored object."""

    project_id = seed_project(db=db, email=f"{uuid4().hex}@example.com")
    seeded = seed_upload(db=db, project_id=project_id, filename=filename, content=content)
    # The seed's commit closed the transaction, so the PUT spans no transaction either.
    storage.put_object(key=seeded.raw_object_key, content=content)
    return project_id, seeded


def _ingest(
    *,
    db: Session,
    seeded: SeededUpload,
    settings: Settings,
    storage: StorageService,
) -> None:
    """Run the ingestion task for one run, unpacking the payload the way the worker will."""

    ingest_source_item(
        db=db,
        source_item_id=seeded.source_item_id,
        ingestion_run_id=seeded.ingestion_run_id,
        raw_object_key=seeded.raw_object_key,
        content_hash=seeded.content_hash,
        storage=storage,
        parser=get_parser(settings=settings),
        settings=settings,
    )


def _read_vectors(*, db: Session, document_id: UUID) -> list[list[float] | None]:
    """Return one document's stored chunk vectors, ordered by chunk position."""

    chunks = db.query(Chunk).filter(Chunk.document_id == document_id).order_by(Chunk.position).all()
    return [chunk.embedding for chunk in chunks]


def test_ingest_end_to_end(
    *,
    db_session: Session,
    settings: Settings,
    storage: StorageService,
) -> None:
    """Check the full happy path into the single indexing transaction."""

    paragraphs = [
        "The first paragraph carries plain prose for the packer.",
        "A second paragraph joins the same pack under the heading.",
        "The third paragraph still fits inside the default target budget.",
    ]
    content = _make_docx(heading="Introduction", paragraphs=paragraphs)
    project_id, seeded = _seed(
        db=db_session,
        storage=storage,
        filename="report.docx",
        content=content,
    )
    _ingest(db=db_session, seeded=seeded, settings=settings, storage=storage)

    run = db_session.get(IngestionRun, seeded.ingestion_run_id)
    assert run is not None
    assert run.status == "indexing"
    assert run.document_id is not None
    assert run.started_at is not None

    document = db_session.get(Document, run.document_id)
    assert document is not None
    parsed = get_parser(settings=settings).parse(filename="report.docx", content=content)
    expected_json = canonical_document_json(document=parsed)
    assert document.content_hash == hashlib.sha256(expected_json.encode("utf-8")).hexdigest()
    assert document.source_item_id == seeded.source_item_id
    assert document.raw_object_key == seeded.raw_object_key
    assert document.normalized_content == {
        "title": parsed.title,
        "elements": [
            {"type": element.type, "text": element.text, "level": element.level}
            for element in parsed.elements
        ],
        "metadata": parsed.metadata,
    }

    chunks = db_session.query(Chunk).filter(Chunk.document_id == document.id).order_by(Chunk.position).all()
    assert len(chunks) == 1
    chunk = chunks[0]
    assert chunk.position == 0
    assert chunk.project_id == project_id
    assert chunk.chunk_type == "text"
    assert chunk.text == "\n\n".join(paragraphs)
    assert chunk.heading_path == ["Introduction"]
    assert chunk.chunk_metadata == {"filename": "report.docx"}
    assert chunk.content_hash == hashlib.sha256(chunk.text.encode("utf-8")).hexdigest()
    assert chunk.embedding is None

    # The project scope is carried down the whole chain: source -> item -> document -> chunk.
    source = db_session.get(Source, seeded.source_id)
    item = db_session.get(SourceItem, seeded.source_item_id)
    assert source is not None
    assert item is not None
    for chain_chunk in chunks:
        assert chain_chunk.project_id == document.project_id
        assert chain_chunk.project_id == item.project_id
        assert chain_chunk.project_id == source.project_id
        assert chain_chunk.project_id == project_id

    # The generated column must equal its own definition, compared inside the database.
    vector_matches = (
        db_session.execute(
            text("SELECT search_vector = to_tsvector('english', text) FROM chunks WHERE document_id = :id"),
            {"id": document.id},
        )
        .scalars()
        .all()
    )
    assert vector_matches == [True]

    events = (
        db_session.query(OutboxEvent).filter(OutboxEvent.event_type == DOCUMENT_READY_FOR_EMBEDDING).all()
    )
    assert len(events) == 1
    assert events[0].aggregate_type == "document"
    assert events[0].aggregate_id == document.id
    assert events[0].payload == {
        "document_id": str(document.id),
        "ingestion_run_id": str(run.id),
    }


def test_duplicate_delivery_no_ops(
    *,
    db_session: Session,
    settings: Settings,
    storage: StorageService,
) -> None:
    """Check that a re-delivered run cannot be re-claimed and writes nothing new."""

    content = _make_docx(heading="Introduction", paragraphs=["One paragraph of prose."])
    _project_id, seeded = _seed(
        db=db_session,
        storage=storage,
        filename="duplicate.docx",
        content=content,
    )
    _ingest(db=db_session, seeded=seeded, settings=settings, storage=storage)
    run = db_session.get(IngestionRun, seeded.ingestion_run_id)
    assert run is not None
    status_before = run.status
    document_id_before = run.document_id
    started_at_before = run.started_at

    _ingest(db=db_session, seeded=seeded, settings=settings, storage=storage)

    db_session.refresh(run)
    assert run.status == status_before
    assert run.document_id == document_id_before
    assert run.started_at == started_at_before
    assert db_session.query(Document).filter(Document.source_item_id == seeded.source_item_id).count() == 1
    events = (
        db_session.query(OutboxEvent).filter(OutboxEvent.event_type == DOCUMENT_READY_FOR_EMBEDDING).all()
    )
    assert len(events) == 1


def test_missing_run_no_ops(
    *,
    db_session: Session,
    settings: Settings,
    storage: StorageService,
) -> None:
    """Check that a delivery naming unknown ids is acked without writing any rows."""

    ingest_source_item(
        db=db_session,
        source_item_id=uuid4(),
        ingestion_run_id=uuid4(),
        raw_object_key="raw/missing/missing/missing/none.pdf",
        content_hash="missing",
        storage=storage,
        parser=get_parser(settings=settings),
        settings=settings,
    )
    assert db_session.query(Document).count() == 0
    assert db_session.query(Chunk).count() == 0
    assert db_session.query(OutboxEvent).count() == 0


def test_garbage_pdf_fails_the_run(
    *,
    db_session: Session,
    settings: Settings,
    storage: StorageService,
) -> None:
    """Check that unparseable bytes fail the run with no document and the error recorded."""

    content = b"these bytes are definitely not a pdf document"
    _project_id, seeded = _seed(
        db=db_session,
        storage=storage,
        filename="broken.pdf",
        content=content,
    )
    _ingest(db=db_session, seeded=seeded, settings=settings, storage=storage)

    run = db_session.get(IngestionRun, seeded.ingestion_run_id)
    assert run is not None
    assert run.status == "failed"
    assert run.error is not None
    assert run.finished_at is not None
    assert db_session.query(Document).filter(Document.source_item_id == seeded.source_item_id).count() == 0


def test_zero_chunk_document_fails_the_run(
    *,
    db_session: Session,
    settings: Settings,
    storage: StorageService,
) -> None:
    """Check that a headings-only document fails the run instead of storing zero chunks."""

    content = _make_docx(heading="Only A Heading", paragraphs=[])
    _project_id, seeded = _seed(
        db=db_session,
        storage=storage,
        filename="headings-only.docx",
        content=content,
    )
    _ingest(db=db_session, seeded=seeded, settings=settings, storage=storage)

    run = db_session.get(IngestionRun, seeded.ingestion_run_id)
    assert run is not None
    assert run.status == "failed"
    assert run.error is not None
    assert "no text content" in run.error
    assert run.finished_at is not None
    assert db_session.query(Document).filter(Document.source_item_id == seeded.source_item_id).count() == 0


def test_identical_version_conflict_is_idempotent_noop(
    *,
    db_session: Session,
    settings: Settings,
    storage: StorageService,
) -> None:
    """Check the unique backstop: a second identical run ends ready on the stored document."""

    content = _make_docx(heading="Introduction", paragraphs=["Shared prose for both runs."])
    _project_id, seeded = _seed(
        db=db_session,
        storage=storage,
        filename="conflict.docx",
        content=content,
    )
    run_b = IngestionRun(id=uuid4(), source_item_id=seeded.source_item_id, status="queued")
    db_session.add(run_b)
    db_session.commit()

    _ingest(db=db_session, seeded=seeded, settings=settings, storage=storage)
    _ingest(
        db=db_session, seeded=replace(seeded, ingestion_run_id=run_b.id), settings=settings, storage=storage
    )

    run_a = db_session.get(IngestionRun, seeded.ingestion_run_id)
    assert run_a is not None
    db_session.refresh(run_b)
    assert run_a.status == "indexing"
    assert run_a.document_id is not None
    assert run_b.status == "ready"
    assert run_b.document_id == run_a.document_id
    documents = db_session.query(Document).filter(Document.source_item_id == seeded.source_item_id).all()
    assert len(documents) == 1
    assert db_session.query(Chunk).filter(Chunk.document_id == run_a.document_id).count() == 1
    events = (
        db_session.query(OutboxEvent).filter(OutboxEvent.event_type == DOCUMENT_READY_FOR_EMBEDDING).all()
    )
    assert len(events) == 1


def test_embed_updates_vectors_and_marks_ready(
    *,
    db_session: Session,
    settings: Settings,
    storage: StorageService,
) -> None:
    """Check that embedding fills every chunk vector and moves the run to ready."""

    paragraphs = ["Alpha prose for the embedding pass.", "Beta prose for the embedding pass."]
    content = _make_docx(heading="Vectors", paragraphs=paragraphs)
    _project_id, seeded = _seed(
        db=db_session,
        storage=storage,
        filename="embed.docx",
        content=content,
    )
    _ingest(db=db_session, seeded=seeded, settings=settings, storage=storage)
    run = db_session.get(IngestionRun, seeded.ingestion_run_id)
    assert run is not None
    assert run.document_id is not None

    embed_document(
        db=db_session,
        document_id=run.document_id,
        ingestion_run_id=run.id,
        embedding_provider=get_embedding_provider(settings=settings),
        settings=settings,
    )

    db_session.refresh(run)
    assert run.status == "ready"
    assert run.finished_at is not None
    chunks = (
        db_session.query(Chunk).filter(Chunk.document_id == run.document_id).order_by(Chunk.position).all()
    )
    assert len(chunks) == 1
    expected = get_embedding_provider(settings=settings).embed_documents(
        texts=[chunk.text for chunk in chunks]
    )
    for chunk, vector in zip(chunks, expected, strict=True):
        assert chunk.embedding is not None
        assert len(chunk.embedding) == STUB_EMBEDDING_DIM
        # The column stores float4, so the round trip only preserves the stub's values approximately.
        assert chunk.embedding == pytest.approx(vector)


def test_embed_is_idempotent_on_redelivery(
    *,
    db_session: Session,
    settings: Settings,
    storage: StorageService,
) -> None:
    """Check that re-delivering the embed task keeps the run ready with identical vectors."""

    content = _make_docx(heading="Stable", paragraphs=["Stable prose embeds identically."])
    _project_id, seeded = _seed(
        db=db_session,
        storage=storage,
        filename="stable.docx",
        content=content,
    )
    _ingest(db=db_session, seeded=seeded, settings=settings, storage=storage)
    run = db_session.get(IngestionRun, seeded.ingestion_run_id)
    assert run is not None
    assert run.document_id is not None
    document_id = run.document_id
    provider = get_embedding_provider(settings=settings)

    embed_document(
        db=db_session,
        document_id=document_id,
        ingestion_run_id=run.id,
        embedding_provider=provider,
        settings=settings,
    )
    vectors_before = _read_vectors(db=db_session, document_id=document_id)

    embed_document(
        db=db_session,
        document_id=document_id,
        ingestion_run_id=run.id,
        embedding_provider=provider,
        settings=settings,
    )

    db_session.refresh(run)
    assert run.status == "ready"
    assert _read_vectors(db=db_session, document_id=document_id) == vectors_before


def test_embed_batches_chunks(*, db_session: Session, storage: StorageService) -> None:
    """Check that embedding splits five chunks into provider calls of the batch size."""

    settings = SettingsForTests(
        chunk_target_tokens=4,
        chunk_max_tokens=8,
        chunk_overlap_tokens=2,
        embedding_batch_size=2,
    )
    paragraphs = [
        "alpha beta gamma",
        "alpha beta delta",
        "alpha beta epsilon",
        "alpha beta zeta",
        "alpha beta eta",
    ]
    content = _make_docx(heading=None, paragraphs=paragraphs)
    _project_id, seeded = _seed(
        db=db_session,
        storage=storage,
        filename="batches.docx",
        content=content,
    )
    _ingest(db=db_session, seeded=seeded, settings=settings, storage=storage)
    run = db_session.get(IngestionRun, seeded.ingestion_run_id)
    assert run is not None
    assert run.document_id is not None
    assert db_session.query(Chunk).filter(Chunk.document_id == run.document_id).count() == 5

    counting = _CountingEmbeddingProvider(inner=get_embedding_provider(settings=settings))
    embed_document(
        db=db_session,
        document_id=run.document_id,
        ingestion_run_id=run.id,
        embedding_provider=counting,
        settings=settings,
    )
    assert counting.call_sizes == [2, 2, 1]
    db_session.refresh(run)
    assert run.status == "ready"


def test_embed_missing_document_no_ops(*, db_session: Session, settings: Settings) -> None:
    """Check that an embed delivery naming an unknown document is acked without rows."""

    embed_document(
        db=db_session,
        document_id=uuid4(),
        ingestion_run_id=uuid4(),
        embedding_provider=get_embedding_provider(settings=settings),
        settings=settings,
    )
    assert db_session.query(IngestionRun).count() == 0
    assert db_session.query(Chunk).count() == 0
