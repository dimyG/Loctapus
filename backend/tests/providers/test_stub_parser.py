"""Tests for the stub parser provider."""

import io
from pathlib import Path

import pytest
from docx import Document

from app.providers import get_parser
from app.providers.base import ParseError
from app.providers.stub_parser import StubParser
from tests.helpers import SettingsForTests

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "fixtures"


def _docx_bytes() -> bytes:
    """Build an in-memory DOCX with headings, paragraphs, and a table."""

    buffer = io.BytesIO()
    document = Document()
    document.add_heading("Chapter", level=1)
    document.add_heading("Section", level=2)
    document.add_paragraph("Body text.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "ignored"
    table.cell(0, 1).text = "also ignored"
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_fixture_yields_one_paragraph_per_page() -> None:
    """Check that the tiny PDF fixture parses into one paragraph element."""

    content = (FIXTURES_DIR / "tiny.pdf").read_bytes()
    document = StubParser().parse(filename="tiny.pdf", content=content)
    assert document.title == "tiny"
    assert document.metadata == {"filename": "tiny.pdf"}
    assert len(document.elements) == 1
    element = document.elements[0]
    assert element.type == "paragraph"
    assert element.level is None
    assert "Stub fixture paragraph." in element.text


def test_docx_headings_and_paragraphs() -> None:
    """Check that DOCX headings, paragraphs, and tables map to the right elements."""

    document = StubParser().parse(filename="sample.docx", content=_docx_bytes())
    assert document.title == "sample"
    assert document.metadata == {"filename": "sample.docx"}
    assert [(element.type, element.level, element.text) for element in document.elements] == [
        ("heading", 1, "Chapter"),
        ("heading", 2, "Section"),
        ("paragraph", None, "Body text."),
    ]


def test_empty_document_raises_parse_error() -> None:
    """Check that unusable content raises ParseError with the user-safe copy."""

    with pytest.raises(ParseError, match="This PDF file could not be read."):
        StubParser().parse(filename="x.pdf", content=b"")


def test_parser_copy_is_user_safe() -> None:
    """Check that every parser failure mode reports its fixed user-facing copy."""

    with pytest.raises(ParseError, match="This file type is not supported."):
        StubParser().parse(filename="x.txt", content=b"text")
    with pytest.raises(ParseError, match="This PDF file could not be read."):
        StubParser().parse(filename="x.pdf", content=b"not a pdf at all")
    with pytest.raises(ParseError, match="This DOCX file could not be read."):
        StubParser().parse(filename="x.docx", content=b"not a docx at all")


def test_unknown_parser_provider_raises() -> None:
    """Check that an unknown parser provider string is rejected."""

    settings = SettingsForTests(parser_provider="nope")
    with pytest.raises(ValueError, match="nope"):
        get_parser(settings=settings)
