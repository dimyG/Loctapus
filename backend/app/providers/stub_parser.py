"""Stub parser extracting plain text from PDF and DOCX documents."""

import io
import re
from pathlib import PurePosixPath

from docx import Document
from pypdf import PdfReader

from app.providers.base import NormalizedDocument, NormalizedElement, ParseError

_HEADING_STYLE = re.compile(r"^Heading (\d+)$")


def _heading_level(*, style_name: str) -> int | None:
    """Return the heading level for a ``Heading N`` style name, otherwise None."""

    match = _HEADING_STYLE.match(style_name)
    return int(match.group(1)) if match else None


class StubParser:
    """Parser that maps each PDF page to a paragraph and DOCX styles to elements."""

    def parse(self, *, filename: str, content: bytes) -> NormalizedDocument:
        """Parse ``content`` into a normalized document chosen by the filename suffix."""

        suffix = PurePosixPath(filename).suffix.lower()
        if suffix == ".pdf":
            elements = self._parse_pdf(content=content)
        elif suffix == ".docx":
            elements = self._parse_docx(content=content)
        else:
            raise ParseError("This file type is not supported.")
        if not elements:
            raise ParseError("No text could be extracted from this file.")
        return NormalizedDocument(
            title=PurePosixPath(filename).stem, elements=elements, metadata={"filename": filename}
        )

    def _parse_pdf(self, *, content: bytes) -> list[NormalizedElement]:
        """Extract one paragraph element per non-empty PDF page."""

        try:
            reader = PdfReader(io.BytesIO(content))
            texts = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            # pypdf raises a mix of exception types for damaged input, so convert them all.
            raise ParseError("This PDF file could not be read.") from exc
        return [NormalizedElement(type="paragraph", text=text.strip()) for text in texts if text.strip()]

    def _parse_docx(self, *, content: bytes) -> list[NormalizedElement]:
        """Extract heading and paragraph elements from a DOCX document, skipping tables."""

        try:
            document = Document(io.BytesIO(content))
        except Exception as exc:
            # python-docx raises zipfile errors for damaged input, so convert them all.
            raise ParseError("This DOCX file could not be read.") from exc
        elements: list[NormalizedElement] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style
            level = _heading_level(style_name=style.name if style is not None else "")
            if level is None:
                elements.append(NormalizedElement(type="paragraph", text=text))
            else:
                elements.append(NormalizedElement(type="heading", text=text, level=level))
        return elements
